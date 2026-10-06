import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Paramètres du site, exportés par parametres.mjs (voir construire.sh) : les classeurs utilisent les mêmes valeurs.
P = json.load(open(os.environ.get("PARAMETRES", os.path.join(os.path.dirname(__file__), "parametres.json"))))
INF = 1e15
F = "Arial"
BLEU = Font(name=F, size=10, color="0000FF")
NOIR = Font(name=F, size=10, color="000000")
VERT = Font(name=F, size=10, color="008000")
GRAS = Font(name=F, size=10, bold=True)
TITRE = Font(name=F, size=14, bold=True)
SOUS = Font(name=F, size=10, italic=True, color="555555")
ENTETE = Font(name=F, size=10, bold=True, color="FFFFFF")
FOND_ENTETE = PatternFill("solid", fgColor="1F3A5F")
JAUNE = PatternFill("solid", fgColor="FFFF00")
FOND_RES = PatternFill("solid", fgColor="E8F0E4")
FIN = Side(style="thin", color="BBBBBB")
BORD = Border(bottom=FIN)

PCT1 = '0.0%;-0.0%;"-"'
PCT2 = '0.00%;-0.00%;"-"'
MD = '#,##0.0" Md€";-#,##0.0" Md€";"-"'
MD2 = '#,##0.00" Md€";-#,##0.00" Md€";"-"'
EUR = '#,##0" €";-#,##0" €";"-"'
EUR2 = '#,##0.00" €";-#,##0.00" €";"-"'
NB = '#,##0;-#,##0;"-"'
NB1 = '#,##0.0;-#,##0.0;"-"'
NB3 = '0.000;-0.000;"-"'

def nouveau():
    wb = Workbook()
    wb.remove(wb.active)
    return wb

def feuille(wb, nom, titre, sous=None, largeurs=None):
    ws = wb.create_sheet(nom)
    ws["A1"] = titre; ws["A1"].font = TITRE
    if sous:
        ws["A2"] = sous; ws["A2"].font = SOUS
    for i, l in enumerate(largeurs or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = l
    ws.sheet_view.showGridLines = False
    return ws

def entete(ws, ligne, valeurs, col=1):
    for i, v in enumerate(valeurs):
        c = ws.cell(ligne, col + i, v); c.font = ENTETE; c.fill = FOND_ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")

def entree(ws, ref, valeur, fmt=None, note=None, cle=True):
    c = ws[ref]; c.value = valeur; c.font = BLEU
    if cle: c.fill = JAUNE
    if fmt: c.number_format = fmt
    if note: c.comment = Comment(note, "Site La grande transmission")
    return c

def formule(ws, ref, f, fmt=None, gras=False, lien=False, res=False):
    c = ws[ref]; c.value = f
    c.font = VERT if lien else (Font(name=F, size=10, bold=True) if gras else NOIR)
    if fmt: c.number_format = fmt
    if res: c.fill = FOND_RES
    return c

def texte(ws, ref, t, gras=False, italique=False, wrap=False):
    c = ws[ref]; c.value = t
    c.font = Font(name=F, size=10, bold=gras, italic=italique, color="555555" if italique else "000000")
    if wrap: c.alignment = Alignment(wrap_text=True, vertical="top")
    return c

def liste(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(ws[ref])

def police_partout(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != F:
                    f = c.font
                    c.font = Font(name=F, size=f.size or 10, bold=f.bold, italic=f.italic, color=f.color)

def lisez_moi(wb, titre, lignes):
    ws = feuille(wb, "Lisez-moi", titre, "La grande transmission : https://bonquentin-collab.github.io/Create-Website/", [110])
    r = 4
    for l in lignes:
        if isinstance(l, tuple):
            texte(ws, f"A{r}", l[0], gras=True)
        else:
            texte(ws, f"A{r}", l, wrap=True)
        r += 1
    r += 1
    texte(ws, f"A{r}", "Chaque chiffre saisi est revérifié à la source : voir la feuille « Vérification » (type, calcul, source, statut).", gras=False); r += 2
    texte(ws, f"A{r}", "Légende des couleurs", gras=True); r += 1
    c = ws[f"A{r}"]; c.value = "Texte bleu sur fond jaune : hypothèse ou donnée d'entrée, à modifier"; c.font = BLEU; c.fill = JAUNE; r += 1
    c = ws[f"A{r}"]; c.value = "Texte bleu sans fond : donnée sourcée (tableau de référence)"; c.font = BLEU; r += 1
    c = ws[f"A{r}"]; c.value = "Texte noir : formule"; c.font = NOIR; r += 1
    c = ws[f"A{r}"]; c.value = "Texte vert : formule qui reprend une valeur d'une autre feuille"; c.font = VERT; r += 1
    c = ws[f"A{r}"]; c.value = "Fond vert pâle : résultat principal"; c.font = NOIR; c.fill = FOND_RES; r += 1
    return ws

def sources(wb, lignes):
    ws = feuille(wb, "Sources", "Sources", None, [42, 80, 70])
    entete(ws, 3, ["Donnée", "Source", "Lien"])
    for i, (d, s, u) in enumerate(lignes, start=4):
        texte(ws, f"A{i}", d, wrap=True); texte(ws, f"B{i}", s, wrap=True)
        if u:
            c = ws[f"C{i}"]; c.value = u; c.hyperlink = u; c.font = Font(name=F, size=10, color="0563C1", underline="single")
    return ws


LIEN = Font(name=F, size=10, color="0563C1", underline="single")
STATUT_FOND = {"Corrigé": PatternFill("solid", fgColor="FCE4D6"), "Estimation, non vérifiable": PatternFill("solid", fgColor="F2F2F2")}

def verification(wb, entrees, date="octobre 2026"):
    """Feuille « Vérification » : chaque chiffre saisi, son type, son calcul, sa source et le résultat de la vérification."""
    ws = feuille(wb, "Vérification", "Vérification des chiffres", f"Chaque chiffre saisi dans ce classeur, revérifié à la source en {date}.", [46, 22, 18, 60, 46, 40, 22, 50])
    entete(ws, 4, ["Donnée", "Valeur", "Type", "Comment c'est calculé", "Source", "Lien", "Statut", "Remarque"])
    for i, e in enumerate(entrees, start=5):
        for col, cle in zip("ABCDEGH", ["libelle", "valeur", "type", "calcul", "source", "statut", "note"]):
            texte(ws, f"{col}{i}", e[cle], wrap=True)
        if e["url"]:
            c = ws[f"F{i}"]; c.value = e["url"]; c.hyperlink = e["url"]; c.font = LIEN; c.alignment = Alignment(wrap_text=True, vertical="top")
        if e["statut"] in STATUT_FOND:
            for col in "ABCDEFGH": ws[f"{col}{i}"].fill = STATUT_FOND[e["statut"]]
    r = 6 + len(entrees)
    texte(ws, f"A{r}", "Types : « Chiffre publié » (repris d'une publication), « Lu sur un graphique », « Calcul du site » (formule à partir de chiffres publiés), « Hypothèse du site », « Loi ou barème ».", italique=True)
    texte(ws, f"A{r+1}", "Statuts : « Confirmé » (vérifié à la source primaire), « Corrigé » (valeur changée lors de cette vérification), « Confirmé (source secondaire) », « Estimation, non vérifiable ».", italique=True)
    return ws

def index_verif(entrees):
    return {e["cle"]: e for e in entrees}

def colonnes_verif(ws, ligne, e, col_type, col_calcul):
    """Écrit le type et le calcul d'un paramètre dans deux colonnes."""
    if not e: return
    texte(ws, f"{col_type}{ligne}", e["type"], wrap=True)
    texte(ws, f"{col_calcul}{ligne}", e["calcul"], wrap=True)
